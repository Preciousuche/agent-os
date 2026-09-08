"""docx skill — load, eligibility, and create→inspect round-trip."""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

from agentos.skills.eligibility import EligibilityContext, check_eligibility
from agentos.skills.loader import SkillLoader

ROOT = Path(__file__).resolve().parents[1]
BUNDLED = ROOT / "src" / "agentos" / "skills" / "bundled"
DOCX_DIR = BUNDLED / "docx"
SCRIPTS = DOCX_DIR / "scripts"


def _spec_to_loader() -> object:
    return SkillLoader(bundled_dir=BUNDLED).get_by_name("docx")


def test_skill_loads() -> None:
    spec = _spec_to_loader()
    assert spec is not None
    assert spec.name == "docx"
    assert spec.metadata is not None
    assert spec.provenance.origin == "clawhub-mit0"
    assert spec.provenance.license == "MIT-0"


def test_eligibility_with_python_present(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "agentos.skills.eligibility.shutil.which",
        lambda name: "/usr/bin/python3" if name in {"python", "python3"} else None,
    )
    spec = _spec_to_loader()
    assert spec is not None
    assert check_eligibility(spec, EligibilityContext.auto())


def test_eligibility_without_python(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "agentos.skills.eligibility.shutil.which",
        lambda name: None,
    )
    spec = _spec_to_loader()
    assert spec is not None
    assert not check_eligibility(spec, EligibilityContext.auto())


def test_create_then_inspect_round_trip(tmp_path: Path) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_docx  # type: ignore[import-not-found]
        import inspect_docx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    spec = {
        "metadata": {"title": "Round-trip", "author": "Tester"},
        "body": [
            {"kind": "heading", "level": 1, "text": "Hello"},
            {"kind": "paragraph", "text": "World."},
            {"kind": "table", "rows": [["A", "B"], ["1", "2"]]},
        ],
    }
    out_path = tmp_path / "out.docx"
    doc = create_docx.build(spec)
    doc.save(str(out_path))
    assert out_path.exists()

    inspected = inspect_docx.inspect(out_path)
    assert inspected["sections"] >= 1
    texts = [p["text"] for p in inspected["paragraphs"]]
    assert "Hello" in texts
    assert "World." in texts
    assert inspected["tables"] and inspected["tables"][0][0] == ["A", "B"]
    assert inspected["has_tracked_changes"] is False


def test_edit_replace_text(tmp_path: Path) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_docx  # type: ignore[import-not-found]
        import edit_docx  # type: ignore[import-not-found]
        import inspect_docx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    src = tmp_path / "src.docx"
    create_docx.build({"body": [{"kind": "paragraph", "text": "Hello {{NAME}}, welcome."}]}).save(
        str(src)
    )

    from docx import Document

    doc = Document(str(src))
    ops = [{"op": "replace_text", "find": "{{NAME}}", "with": "Wei"}]
    edit_docx.apply_ops(doc, ops)
    out = tmp_path / "out.docx"
    doc.save(str(out))

    inspected = inspect_docx.inspect(out)
    text = " ".join(p["text"] for p in inspected["paragraphs"])
    assert "{{NAME}}" not in text
    assert "Wei" in text


def _load_edit_docx() -> object:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import edit_docx  # type: ignore[import-not-found]

        return edit_docx
    finally:
        sys.path.pop(0)


def test_replace_text_inside_a_single_run_only_touches_that_run() -> None:
    """A match fully inside one run must not disturb any other run's text
    or formatting (#1447)."""
    from docx import Document

    edit_docx = _load_edit_docx()
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Hello ")
    bold = para.add_run("world")
    bold.bold = True
    para.add_run(" and more")

    changed = edit_docx.apply_ops(doc, [{"op": "replace_text", "find": "Hello", "with": "Hi"}])

    assert changed == 1
    texts = [run.text for run in para.runs]
    assert texts == ["Hi ", "world", " and more"]
    assert para.runs[1].bold is True


def test_replace_text_spanning_two_runs_preserves_untouched_runs() -> None:
    """A match spanning the boundary between run 0 and run 1 must land in
    run 0 (the run owning the match's first character); run 1 keeps only
    its unmatched characters, and run 2 -- which the match never reached --
    must stay completely untouched, formatting included (#1447)."""
    from docx import Document

    edit_docx = _load_edit_docx()
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Hello ")
    bold = para.add_run("world")
    bold.bold = True
    tail = para.add_run(" and more")

    changed = edit_docx.apply_ops(doc, [{"op": "replace_text", "find": "o world", "with": "!"}])

    assert changed == 1
    assert [run.text for run in para.runs] == ["Hell!", "", " and more"]
    # The run that only lost matched characters still exists, with its
    # formatting intact even though its text is now empty.
    assert bold.bold is True
    assert tail.text == " and more"


def test_replace_text_multiple_matches_in_one_paragraph() -> None:
    """Every non-overlapping occurrence is replaced, matching str.replace's
    own left-to-right, non-overlapping semantics (#1447)."""
    from docx import Document

    edit_docx = _load_edit_docx()
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("cat sat on the cat mat")

    changed = edit_docx.apply_ops(doc, [{"op": "replace_text", "find": "cat", "with": "dog"}])

    assert changed == 1
    assert para.runs[0].text == "dog sat on the dog mat"


def test_replace_text_find_not_present_leaves_paragraph_byte_identical() -> None:
    """A find that never matches must not touch any run at all (#1447)."""
    from docx import Document

    edit_docx = _load_edit_docx()
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Hello ")
    bold = para.add_run("world")
    bold.bold = True

    changed = edit_docx.apply_ops(doc, [{"op": "replace_text", "find": "missing", "with": "x"}])

    assert changed == 0
    assert [run.text for run in para.runs] == ["Hello ", "world"]
    assert para.runs[1].bold is True


def test_replace_run_is_unaffected_by_the_replace_text_fix() -> None:
    from docx import Document

    edit_docx = _load_edit_docx()
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("one")
    para.add_run("two")

    changed = edit_docx.apply_ops(doc, [{"op": "replace_run", "para": 0, "run": 1, "text": "TWO"}])

    assert changed == 1
    assert [run.text for run in para.runs] == ["one", "TWO"]


def test_inspect_cli_outputs_json(tmp_path: Path) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_docx  # type: ignore[import-not-found]
        import inspect_docx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    src = tmp_path / "src.docx"
    create_docx.build({"body": [{"kind": "paragraph", "text": "x"}]}).save(str(src))

    payload = inspect_docx.inspect(src)
    encoded = json.dumps(payload, ensure_ascii=False)
    assert "paragraphs" in encoded
    assert "tables" in encoded


def test_inspect_docx_creates_parent_directory(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_docx  # type: ignore[import-not-found]
        import inspect_docx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    src = tmp_path / "src.docx"
    create_docx.build({"body": [{"kind": "paragraph", "text": "x"}]}).save(str(src))

    out = tmp_path / "nested" / "dir" / "out.json"
    monkeypatch.setattr(sys, "argv", ["inspect_docx.py", str(src), "--out", str(out)])
    assert inspect_docx.main() == 0
    assert out.is_file()
